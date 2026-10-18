from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
doc=Document()
# Remove the default template title rule and any inherited paragraph borders.
for node in list(doc.styles.element.iter(qn('w:pBdr'))):
    node.getparent().remove(node)
section=doc.sections[0]
section.page_width=Inches(8.27)
section.page_height=Inches(11.69)
section.top_margin=section.bottom_margin=Inches(.72)
section.left_margin=section.right_margin=Inches(.73)
section.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name]
    style.font.name='Arial'
    style.font.color.rgb=RGBColor(0,0,0)
    rpr=style.element.get_or_add_rPr()
    for c in list(rpr):
        if c.tag==qn('w:color'):
            for a in ['themeColor','themeShade','themeTint']:
                c.attrib.pop(qn('w:'+a),None)
normal=doc.styles['Normal']
normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.1
normal.paragraph_format.widow_control=True
for name,size in [('Title',23),('Heading 1',16),('Heading 2',12.5),('Heading 3',11.5)]:
    doc.styles[name].font.size=Pt(size)
    if name!='Title':doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles[name].paragraph_format.space_before=Pt(12 if name!='Title' else 0)
    doc.styles[name].paragraph_format.space_after=Pt(7)
doc.styles['Subtitle'].font.size=Pt(12)
doc.styles['Subtitle'].font.italic=False
doc.core_properties.title='Usulan revisi cerita dan pilihan pemain Wildroot'
doc.core_properties.subject='Delapan usulan untuk keputusan Stefanus Terang'
doc.core_properties.author='Codex dengan tinjauan MiniMax M3'

def add_text(text,style=None):
    p=doc.add_paragraph(style=style)
    chunks=re.split(r'(\*\*.*?\*\*)',text)
    for ch in chunks:
        r=p.add_run(ch[2:-2] if ch.startswith('**') else ch)
        if ch.startswith('**'):r.bold=True
    return p

def approval(rid):
    p=add_text('Keputusan '+rid+'    ACC / ACC DENGAN CATATAN / TOLAK / TUNDA')
    p.paragraph_format.space_before=Pt(10)
    p.paragraph_format.keep_with_next=True
    for r in p.runs:r.font.size=Pt(9);r.bold=True
    p=add_text('Catatan pemilik __________________________________________________')
    for r in p.runs:r.font.size=Pt(9)

def table(lines):
    rows=[[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
    rows=[row for row in rows if not all(re.fullmatch(r'[-: ]+',v) for v in row)]
    count=len(rows[0]);t=doc.add_table(rows=0,cols=count)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    usable=6.81
    if rows[0][0]=='ID':ratios=[.08,.75,.17]
    elif rows[0][0]=='Aspek':ratios=[.14,.285,.285,.29]
    elif count==4:ratios=[.16,.28,.28,.28]
    elif count==3:ratios=[.28,.36,.36]
    else:ratios=[.62,.38]
    for col,ratio in zip(t.columns,ratios):col.width=Inches(usable*ratio)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        trpr=t.rows[-1]._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit');trpr.append(no)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,(cell,txt) in enumerate(zip(cells,row)):
            cell.width=Inches(usable*ratios[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=txt
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'33455D' if i==0 else ('F1F3F5' if i%2==0 else 'FFFFFF'));tcpr.append(shade)
            mar=OxmlElement('w:tcMar')
            for side,value in [('top',60),('bottom',60),('left',95),('right',95)]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');mar.append(e)
            tcpr.append(mar)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.04
                if i==0 or (rows[0][0]=='ID' and j!=1):p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name='Arial';r.font.size=Pt(9.3);r.bold=i==0
                    r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

lines=(ROOT/'proposal.md').read_text().splitlines()
i=0; current=None; cover=True
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('|'):
        group=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            group.append(lines[i]);i+=1
        table(group);continue
    if line.startswith('# '):add_text(line[2:],'Title');i+=1;continue
    if line.startswith('## '):
        title=line[3:]
        m=re.match(r'R([1-8]) ',title)
        if m:
            if current:approval(current)
            current='R'+m.group(1)
            p=add_text(title,'Heading 1');p.paragraph_format.page_break_before=True
        else:
            if current:approval(current);current=None
            p=add_text(title,'Heading 1')
            if title in ['Hasil pembahasan Codex dan MiniMax M3']:
                p.paragraph_format.page_break_before=True
        i+=1;continue
    if line.startswith('### '):add_text(line[4:],'Heading 2');i+=1;continue
    if re.match(r'^\d+\. ',line):
        p=add_text(line);p.paragraph_format.left_indent=Inches(.17);p.paragraph_format.first_line_indent=Inches(-.17);i+=1;continue
    group=[line];i+=1
    while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|')) and not re.match(r'^\d+\. ',lines[i]):
        group.append(lines[i].strip());i+=1
    add_text(' '.join(group))
if current:approval(current)
for node in list(doc.element.iter(qn('w:pBdr'))):
    node.getparent().remove(node)
foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('Wildroot usulan revisi   ');r.font.name='Arial';r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
out=ROOT/'Wildroot_Usulan_Revisi_8_Poin_v1.docx'
doc.save(out)
print(out)
print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables))
